#!/usr/bin/env python3
"""
Generate a formatted Excel risk report.

Usage:
  python scripts/generate_report.py                      # outputs/risk_report_YYYYMMDD.xlsx
  python scripts/generate_report.py --notional 2000000
  python scripts/generate_report.py --dollar-neutral
"""

import sys
from pathlib import Path
from datetime import datetime, timedelta

sys.path.insert(0, str(Path(__file__).parent.parent))

import json
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from src.data_loader import load_pairs, load_ticker_map, fetch_prices, fetch_index
from src.position_sizing import (
    compute_betas, beta_neutral_sizes, dollar_neutral_sizes, portfolio_summary
)
from src.factor_attribution import build_factor_table, attribution_summary
from src.sensitivity import load_sensitivities, run_scenarios, SCENARIOS
from src.catalyst_tracker import _load as load_catalysts, summary as catalyst_summary

# ── Palette ─────────────────────────────────────────────────────────────────
NAVY      = "1B2A4A"
TEAL      = "1D6FA4"
LIGHT_BLU = "DCE9F5"
WHITE     = "FFFFFF"
GREEN     = "1A7C4A"
RED_      = "C0392B"
GOLD      = "B8860B"
GREY_HDR  = "3A3A3A"
ALT_ROW   = "F4F8FC"

def hdr_fill(hex_color):  return PatternFill("solid", fgColor=hex_color)
def font(bold=False, color=WHITE, size=10):
    return Font(bold=bold, color=color, name="Calibri", size=size)
def center(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def left():   return Alignment(horizontal="left",   vertical="center")
def right():  return Alignment(horizontal="right",  vertical="center")
def thin_border():
    s = Side(style="thin", color="B0BEC5")
    return Border(left=s, right=s, top=s, bottom=s)

def write_header_row(ws, row_num, values, bg=NAVY, fg=WHITE, row_height=22):
    ws.row_dimensions[row_num].height = row_height
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row_num, column=col, value=val)
        c.fill      = hdr_fill(bg)
        c.font      = font(bold=True, color=fg)
        c.alignment = center()
        c.border    = thin_border()

def write_data_row(ws, row_num, values, formats=None, alt=False):
    bg = ALT_ROW if alt else WHITE
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row_num, column=col, value=val)
        c.fill      = hdr_fill(bg)
        c.font      = Font(name="Calibri", size=9, color="1A1A1A")
        c.alignment = right() if isinstance(val, (int, float)) else left()
        c.border    = thin_border()
        if formats and col <= len(formats) and formats[col-1]:
            c.number_format = formats[col-1]

def set_col_widths(ws, widths):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

def sheet_title(ws, title, subtitle, col_span):
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 16
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_span)
    c1 = ws.cell(row=1, column=1, value=title)
    c1.fill      = hdr_fill(NAVY)
    c1.font      = Font(bold=True, color=WHITE, size=13, name="Calibri")
    c1.alignment = center()
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.fill      = hdr_fill(TEAL)
    c2.font      = Font(bold=False, color=WHITE, size=9, name="Calibri")
    c2.alignment = center()

# ─────────────────────────────────────────────────────────────────────────────
# CLI args
# ─────────────────────────────────────────────────────────────────────────────

args              = sys.argv[1:]
notional_per_pair = 1_000_000
dollar_neutral    = "--dollar-neutral" in args
for i, a in enumerate(args):
    if a == "--notional" and i + 1 < len(args):
        notional_per_pair = float(args[i + 1])

today      = datetime.today().strftime("%Y-%m-%d")
today_slug = datetime.today().strftime("%Y%m%d")

# ─────────────────────────────────────────────────────────────────────────────
# Load data
# ─────────────────────────────────────────────────────────────────────────────

print("Loading pairs and tickers…")
pairs      = load_pairs()
ticker_map = load_ticker_map()

all_names   = []
for _, row in pairs.iterrows():
    all_names += row["long_list"] + row["short_list"]
all_names   = list(dict.fromkeys(all_names))
all_tickers = [ticker_map.get(n) for n in all_names]
valid_tickers = [t for t in all_tickers if t]

start_date = (datetime.today() - timedelta(days=400)).strftime("%Y-%m-%d")
print(f"Fetching prices for {len(valid_tickers)} tickers (since {start_date})…")

prices = fetch_prices(valid_tickers, start=start_date, end=today)
index  = fetch_index("000300", start=start_date, end=today)

if not prices.empty and not index.empty:
    from src.position_sizing import compute_betas
    betas = compute_betas(prices, index)
    print(f"  Computed betas from {len(index)} trading days of CSI-300 data")
else:
    betas = {t: 1.0 for t in valid_tickers}
    print("  [warn] No market data — beta = 1.0 for all (run locally with akshare)")

if dollar_neutral:
    positions = dollar_neutral_sizes(pairs, ticker_map, notional_per_pair)
    positions["beta"] = positions.get("ticker", pd.Series()).map(lambda t: betas.get(t, 1.0) if t else 1.0)
    positions["beta_contribution"] = positions["notional"] * positions["beta"]
    mode = "Dollar-Neutral"
else:
    positions = beta_neutral_sizes(pairs, ticker_map, betas, notional_per_pair)
    mode = "Beta-Neutral"

factor_df     = build_factor_table(positions, betas, prices if not prices.empty else None)
sensitivity   = load_sensitivities()
scenarios_df  = run_scenarios(positions, sensitivity)
catalysts     = load_catalysts()
summ          = portfolio_summary(positions)
gross         = summ["gross_notional"]

# ─────────────────────────────────────────────────────────────────────────────
# Build workbook
# ─────────────────────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
wb.remove(wb.active)   # remove default sheet

subtitle_base = f"As of {today}  |  {mode}  |  Notional/pair: {notional_per_pair:,.0f} CNY  |  Gross: {gross:,.0f} CNY"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 1 — POSITIONS
# ════════════════════════════════════════════════════════════════════════════

ws1 = wb.create_sheet("1. Positions")
sheet_title(ws1, "📊  POSITION BOOK — 2026 Long-Short Pairs", subtitle_base, 9)

hdrs = ["Sector 板块", "Side", "Name 公司", "Ticker", "Beta (β)",
        "Notional (CNY)", "% of Gross", "Investment Thesis", "Sector Catalysts"]
write_header_row(ws1, 3, hdrs, bg=GREY_HDR)

# Build thesis + catalyst lookup from pairs
thesis_map   = {}
catalyst_map = {}
for _, row in pairs.iterrows():
    sector = row["板块"]
    thesis_map[sector]   = row.get("2026边际变化（投资逻辑）", "")
    catalyst_map[sector] = row.get("后续催化剂", "")

row_num = 4
alt = False
for _, pos in positions.sort_values(["sector", "side"], ascending=[True, False]).iterrows():
    side_color = GREEN if pos["side"] == "Long" else RED_
    bet        = pos.get("beta") or betas.get(pos["ticker"], 1.0)
    notional   = pos["notional"]
    pct_gross  = abs(notional) / gross if gross else 0

    vals = [
        pos["sector"],
        pos["side"],
        pos["name"],
        pos["ticker"] or "VERIFY",
        round(float(bet), 2),
        abs(notional),
        pct_gross,
        thesis_map.get(pos["sector"], ""),
        catalyst_map.get(pos["sector"], ""),
    ]
    fmts = [None, None, None, None, "0.00", '#,##0', '0.0%', None, None]
    write_data_row(ws1, row_num, vals, formats=fmts, alt=alt)

    # colour Side cell
    side_cell       = ws1.cell(row=row_num, column=2)
    side_cell.fill  = hdr_fill(GREEN if pos["side"] == "Long" else RED_)
    side_cell.font  = Font(bold=True, color=WHITE, size=9, name="Calibri")
    side_cell.alignment = center()

    row_num += 1
    alt = not alt

# Summary row
ws1.row_dimensions[row_num].height = 18
summ_vals = ["TOTAL / NET", "", "", "", "",
             summ["net_notional"], summ["net_notional"] / gross if gross else 0,
             f"L/S ratio: {summ['l_s_ratio']:.3f}", ""]
for col, val in enumerate(summ_vals, 1):
    c = ws1.cell(row=row_num, column=col, value=val)
    c.fill  = hdr_fill(NAVY)
    c.font  = Font(bold=True, color=WHITE, size=9, name="Calibri")
    c.alignment = center()
    c.border = thin_border()
    if col == 6: c.number_format = '#,##0'
    if col == 7: c.number_format = '0.0%'

set_col_widths(ws1, [12, 7, 12, 8, 7, 16, 10, 42, 32])
ws1.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 2 — SCENARIO ANALYSIS
# ════════════════════════════════════════════════════════════════════════════

ws2 = wb.create_sheet("2. Scenario Analysis")
sheet_title(ws2, "📐  SCENARIO ANALYSIS — Greeks-like Factor Sensitivities",
            subtitle_base + "  |  Elasticities in data/factor_sensitivities.json", 7)

hdrs2 = ["Scenario", "Factor", "Shock", "Long P&L (CNY)", "Short P&L (CNY)", "Net P&L (CNY)", "Net P&L (%)"]
write_header_row(ws2, 3, hdrs2, bg=GREY_HDR)

row_num = 4
for i, (label, (factor, shock)) in enumerate(SCENARIOS.items()):
    row_result = scenarios_df[scenarios_df["Scenario"] == label].iloc[0]
    net_pnl = row_result["Net P&L"]
    net_pct = net_pnl / gross if gross else 0

    vals = [
        label,
        factor,
        f"{shock:+.0%}",
        row_result["Long P&L"],
        row_result["Short P&L"],
        net_pnl,
        net_pct,
    ]
    fmts = [None, None, None, '#,##0', '#,##0', '#,##0', '0.00%']
    write_data_row(ws2, row_num, vals, formats=fmts, alt=(i % 2 == 1))

    # Colour Net P&L cell
    net_cell      = ws2.cell(row=row_num, column=6)
    net_pct_cell  = ws2.cell(row=row_num, column=7)
    color = GREEN if net_pnl >= 0 else RED_
    for cell in (net_cell, net_pct_cell):
        cell.font = Font(bold=True, color=color, size=9, name="Calibri")

    row_num += 1

# Factor legend
row_num += 1
ws2.cell(row=row_num, column=1, value="Factor Definitions:").font = Font(bold=True, size=9, name="Calibri")
factor_defs = {
    "ai_capex":     "Global AI capital expenditure (proxy: NVDA revenue YoY %)",
    "ccl_price":    "Domestic CCL (覆铜板) laminate ASP index",
    "copper_price": "LME copper spot price (USD/t)",
    "silicon_price":"Polysilicon spot price (CNY/kg)",
    "battery_price":"Lithium battery cell price (CNY/Wh)",
    "natgas_price": "TTF / Henry Hub natural gas (USD/MMBtu)",
    "robot_demand": "Humanoid robot production proxy (Tesla Optimus cumulative units)",
}
for i, (f, d) in enumerate(factor_defs.items()):
    ws2.cell(row=row_num + 1 + i, column=1, value=f)
    ws2.cell(row=row_num + 1 + i, column=2, value=d).font = Font(size=8, color="555555", name="Calibri")

set_col_widths(ws2, [22, 14, 8, 16, 16, 16, 12])
ws2.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 3 — CATALYST TRACKER
# ════════════════════════════════════════════════════════════════════════════

ws3 = wb.create_sheet("3. Catalyst Tracker")
c_sum = catalyst_summary()
subtitle3 = (f"As of {today}  |  "
             f"Total: {c_sum['total']}  |  "
             f"✅ Fired: {c_sum['fired']}  |  "
             f"⏳ Pending: {c_sum['pending']}  |  "
             f"❌ Missed: {c_sum['missed']}")
sheet_title(ws3, "📅  CATALYST TRACKER — 2026 Long-Short Pairs", subtitle3, 6)

STATUS_COLORS = {"pending": GOLD, "fired": GREEN, "missed": RED_}
STATUS_LABELS = {"pending": "⏳ Pending", "fired": "✅ Fired", "missed": "❌ Missed"}

write_header_row(ws3, 3, ["Sector 板块", "Catalyst 催化剂", "Status", "Date", "Notes 备注", "Pairs Affected"], bg=GREY_HDR)

# Build a reverse map: sector → pair names
sector_names = {}
for _, row in pairs.iterrows():
    sec = row["板块"]
    names = row["long_list"] + row["short_list"]
    sector_names[sec] = " | ".join(names)

row_num = 4
alt = False
for sector, items in catalysts.items():
    if sector.startswith("_"):
        continue
    for item in items:
        st     = item["status"]
        vals   = [
            sector,
            item["text"],
            STATUS_LABELS[st],
            item.get("fired_date") or "—",
            item.get("notes") or "",
            sector_names.get(sector, ""),
        ]
        write_data_row(ws3, row_num, vals, alt=alt)

        # Colour status cell
        st_cell       = ws3.cell(row=row_num, column=3)
        st_cell.fill  = hdr_fill(STATUS_COLORS[st])
        st_cell.font  = Font(bold=True, color=WHITE, size=9, name="Calibri")
        st_cell.alignment = center()

        row_num += 1
        alt = not alt

set_col_widths(ws3, [12, 26, 13, 12, 32, 40])
ws3.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 4 — FACTOR ATTRIBUTION
# ════════════════════════════════════════════════════════════════════════════

ws4 = wb.create_sheet("4. Factor Attribution")
sheet_title(ws4, "🔬  FACTOR ATTRIBUTION — Market Beta / Momentum / Size",
            subtitle_base, 7)

# 4a — Summary table
write_header_row(ws4, 3, ["Factor", "Long Exposure", "Short Exposure", "Net Exposure",
                           "Interpretation"], bg=GREY_HDR)

agg = attribution_summary(factor_df)
factor_notes = {
    "beta":     "Market beta vs CSI 300. Net ~0 = beta-neutral.",
    "momentum": "+ve = long leg has stronger price momentum. Momentum crowding risk if > +0.5.",
    "size_z":   "+ve = long leg larger cap. -ve = short is larger cap (typical A-share short squeeze risk).",
}
row_num = 4
for i, factor in enumerate(agg.index):
    l_exp  = agg.loc[factor, "Long"]  if "Long"  in agg.columns else np.nan
    s_exp  = agg.loc[factor, "Short"] if "Short" in agg.columns else np.nan
    net_exp = agg.loc[factor, "Net"]  if "Net"   in agg.columns else np.nan
    vals = [factor, l_exp, s_exp, net_exp, factor_notes.get(factor, "")]
    fmts = [None, "+0.000", "+0.000", "+0.000", None]
    write_data_row(ws4, row_num, vals, formats=fmts, alt=(i % 2 == 1))
    row_num += 1

# 4b — Per-sector beta
row_num += 2
ws4.cell(row=row_num, column=1, value="BETA BY SECTOR").font = Font(bold=True, size=10, color=NAVY, name="Calibri")
row_num += 1
write_header_row(ws4, row_num, ["Sector", "Long Beta", "Short Beta", "Net Beta", "Gross Notional (CNY)"], bg=TEAL)
row_num += 1

sector_data = (
    factor_df.groupby(["sector", "side"])
    .apply(
        lambda g: float(np.average(g["beta"].fillna(1.0), weights=g["notional"].abs())),
        include_groups=False,
    )
    .unstack("side")
)
sector_notional = positions.groupby("sector")["notional"].apply(lambda x: x.abs().sum())

for i, sec in enumerate(sector_data.index):
    l_beta = sector_data.loc[sec, "Long"]  if "Long"  in sector_data.columns else 1.0
    s_beta = sector_data.loc[sec, "Short"] if "Short" in sector_data.columns else 1.0
    net_b  = l_beta - s_beta
    gn     = sector_notional.get(sec, 0)
    vals   = [sec, l_beta, s_beta, net_b, gn]
    fmts   = [None, "+0.00", "+0.00", "+0.00", '#,##0']
    write_data_row(ws4, row_num, vals, formats=fmts, alt=(i % 2 == 1))

    net_cell = ws4.cell(row=row_num, column=4)
    net_cell.font = Font(
        bold=True, size=9, name="Calibri",
        color=GREEN if net_b >= -0.05 else RED_
    )
    row_num += 1

set_col_widths(ws4, [16, 15, 15, 14, 20, 42])
ws4.freeze_panes = "A4"

# ─────────────────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────────────────

out_dir = Path(__file__).parent.parent / "outputs"
out_dir.mkdir(exist_ok=True)
out_path = out_dir / f"risk_report_{today_slug}.xlsx"
wb.save(out_path)
print(f"\n✅  Report saved → {out_path.relative_to(Path(__file__).parent.parent)}")
print(f"   Sheets: {[ws.title for ws in wb.worksheets]}")
